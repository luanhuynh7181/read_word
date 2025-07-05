from tkinter import CENTER
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from typing import List
from models.exercise.exercise import Exercise
from docx.oxml import OxmlElement
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT 
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.shared import Inches
from write_doc.utils_write_doc import add_style_text, add_style_paragraph, create_table, fix_column_widths, add_style_cell

def add_1_header(doc: Document, exercises: List[Exercise], de_so: int):  # type: ignore
    add_header_de_so(doc, de_so)
    table =create_table(doc, len(exercises) + 1, 4)
    fix_column_widths(table, [1, 3,1 , 1])
    add_title_table(table)
    add_data_table(table, exercises, de_so)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                 
    
def add_header_de_so(doc: Document, de_so: int):  # type: ignore
    para = doc.add_heading(f"ĐỀ SỐ {de_so}", level=1)
    doc.add_paragraph()
    style = {
        'font_name': 'Arial',
        'font_size': 21,
        'rgb_color': RGBColor(192, 0, 0),
        'alignment': WD_PARAGRAPH_ALIGNMENT.CENTER,
        'line_spacing': 30
    }
    add_style_paragraph(para, style)

def add_title_table(table):
    style = {
        'rgb_color': RGBColor(0, 111, 192),
        'alignment': WD_PARAGRAPH_ALIGNMENT.CENTER,
        'bold': True,
        'vertical_alignment': WD_CELL_VERTICAL_ALIGNMENT.CENTER,
        'space_before': 15,
        'space_after': 15,
    }

    add_style_cell(table.cell(0, 0), 'Bài', style)
    add_style_cell(table.cell(0, 1), 'Tên bài', style)
    add_style_cell(table.cell(0, 2), 'File', style)
    add_style_cell(table.cell(0, 3), 'Điểm', style)

def add_data_table(table, exercises: List[Exercise], de_so: int):
    style_1 = {
        'font_name': 'Times New Roman',
        'font_size': 14,
        'rgb_color': RGBColor(0, 111, 192),
        'alignment': WD_PARAGRAPH_ALIGNMENT.CENTER,
        'bold': True ,
        'vertical_alignment': WD_CELL_VERTICAL_ALIGNMENT.CENTER,
        'space_before': 15,
        'space_after': 15,
    }
    style_2 = style_1.copy()
    style_2['alignment'] = WD_PARAGRAPH_ALIGNMENT.LEFT
    style_2['rgb_color'] = RGBColor(0, 0, 0)
    
    style_3_4 = style_1.copy()
    style_3_4['rgb_color'] = RGBColor(0, 0, 0)
    style_3_4['bold'] = False
    
    for i, exercise in enumerate(exercises):
        add_style_cell(table.cell(i + 1, 0), "Bài " + str(i + 1), style_1)
        add_style_cell(table.cell(i + 1, 1), exercise.title, style_2)
        add_style_cell(table.cell(i + 1, 2), "Ex_" + str(de_so) + "_" + str(i + 1), style_3_4)
        add_style_cell(table.cell(i + 1, 3), "0", style_3_4)

